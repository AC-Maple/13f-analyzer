"""
export_bloomberg_template.py -- Generates an Excel workbook with live
Bloomberg formulas for every distinct underlying security in a classified
13F, for the user to open on a Bloomberg-connected machine, refresh, and
hand back to import_bloomberg_data.py.

DO NOT run recalc.py / LibreOffice on the output of this script. =BDP()
is not a native Excel function -- it is implemented by Bloomberg's own
Excel Add-in, which requires a running Terminal session. LibreOffice has
no Bloomberg connectivity and will write #NAME? into every formula cell,
permanently destroying them. This file is only usable opened for real, on
a machine with the Bloomberg Terminal running.

Deduplicated by CUSIP, not by Security ID or raw row. Price and ADV are
properties of the underlying security -- CYTK's 4 filing rows (common,
2 calls, 1 put) all need exactly the same market data, so pulling it
once per CUSIP avoids 4x redundant Bloomberg queries per name.

Field mnemonics verified live, 2026-09-02:
  - /cusip/{cusip} Equity as a BDP identifier: confirmed against NVDA
    (CUSIP 67066G104) -- returned correct PX_LAST.
  - PARSEKYABLE_DES (field DS587): confirmed against the same CUSIP --
    returned "NVDA US Equity", the full resolved identifier, not just
    a bare ticker. Included below as a fallback ticker-resolution
    source for security_master.py (edgartools' offline CUSIP->ticker
    map is primary; this covers whatever it misses -- recent ticker
    changes, thin small-caps -- and costs nothing extra since it rides
    the same workbook already being pulled for price/ADV).
  - PX_LAST, EQY_SH_OUT, CUR_MKT_CAP, VOLUME_AVG_20D, VOLUME_AVG_3M:
    user-confirmed correct against a live Bloomberg session.
Not yet verified: behavior on a CUSIP Bloomberg doesn't cover (OTC
warrants, delisted names) -- expect a Bloomberg error string there,
handled by import_bloomberg_data.py same as any other field.

GICS_INDUSTRY_NAME / GICS_SUB_INDUSTRY_NAME: added for sector-level
concentration (analyze.py's compute_sector_concentration). Same
/cusip/ identifier, same BDP() pattern as every field above. The
mnemonics themselves are live-confirmed against a real Bloomberg
Terminal refresh, 2026-09-08 -- see SKILL.md's "Sector classification
field verification (GICS, live-tested)" section for the actual
results (CYTK/BSX/SBUX, full GICS level 1-4 breakdown, and the
GICS-vs-BICS decision). What that test did NOT cover: whether these
two fields resolve correctly pulled through THIS script's per-CUSIP
/cusip/{cusip} Equity template pattern at scale, as opposed to the
standalone test workbook's ticker-based pull -- that real refresh is
still owed; see SKILL.md's "Sector concentration wired into the
pipeline" section. Requested for every security (not gated by
needs_adv) since sector classification applies to warrants and
options' underlyings too, not just ADV-modeled common stock -- same
reasoning as PARSEKYABLE_DES below.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DEFAULT_INPUT_FILE = "data/classified_rows.json"
DEFAULT_OUTPUT_FILE = "bloomberg_template.xlsx"

# Instrument classes with no meaningful traded volume of their own --
# still priced (for the implied-price cross-check) but ADV is left blank
# rather than requesting a field that's structurally not applicable.
NO_ADV_CLASSES = {"WARRANT"}

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
INPUT_FONT = Font(name="Arial", color="0000FF")     # blue -- hardcoded input (CUSIP)
FORMULA_FONT = Font(name="Arial", color="000000")   # black -- Bloomberg formula
LABEL_FONT = Font(name="Arial", italic=True, size=9, color="666666")

HEADER_ROW = 4
HEADERS = [
    "Issuer", "CUSIP", "Instrument Classes Covered",
    "PX_LAST", "EQY_SH_OUT (mm)", "CUR_MKT_CAP (mm)",
    "VOLUME_AVG_20D", "VOLUME_AVG_3M", "PARSEKYABLE_DES",
    "GICS_INDUSTRY_NAME", "GICS_SUB_INDUSTRY_NAME",
]


def export_template(rows, output_file):
    """Write a Bloomberg BDP() workbook for these classified rows.

    Same CUSIP collapse and formulas as the original module-level script.
    Does not read or write classified_rows.json / market_data.json.
    """
    by_cusip = {}
    for r in rows:
        cusip = r["cusip"]
        entry = by_cusip.setdefault(cusip, {
            "cusip": cusip,
            "nameOfIssuer": r["nameOfIssuer"],
            "classes": set(),
        })
        entry["classes"].add(r["instrumentClass"])

    securities = sorted(by_cusip.values(), key=lambda e: e["nameOfIssuer"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Bloomberg Pull"

    ws["A1"] = "13F MARKET DATA REQUEST -- BLOOMBERG EXCEL ADD-IN"
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws["A2"] = (
        "Open on a machine with the Bloomberg Terminal running. Formulas "
        "below use the /cusip/ identifier so no separate ticker mapping is "
        "needed. Refresh (Bloomberg ribbon > Refresh, or Ctrl+Alt+F9), save, "
        "and return this file -- do not edit columns C onward by hand."
    )
    ws["A2"].font = LABEL_FONT
    ws.merge_cells("A2:H2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for i, sec in enumerate(securities):
        row = HEADER_ROW + 1 + i
        cusip = sec["cusip"]
        classes = sec["classes"]
        needs_adv = not classes.issubset(NO_ADV_CLASSES)

        ws.cell(row=row, column=1, value=sec["nameOfIssuer"]).font = FORMULA_FONT
        cusip_cell = ws.cell(row=row, column=2, value=cusip)
        cusip_cell.font = INPUT_FONT
        ws.cell(row=row, column=3, value=", ".join(sorted(classes))).font = LABEL_FONT

        # Bloomberg identifier: CUSIP-keyed, not ticker-keyed, so this runs
        # without waiting on a separate CUSIP->ticker resolution step.
        # Verified 2026-09-02 against NVDA (CUSIP 67066G104) -- resolves.
        bbg_id = f'"/cusip/{cusip} Equity"'

        ws.cell(row=row, column=4,
                 value=f'=BDP({bbg_id},"PX_LAST")').font = FORMULA_FONT
        ws.cell(row=row, column=5,
                 value=f'=BDP({bbg_id},"EQY_SH_OUT")').font = FORMULA_FONT
        ws.cell(row=row, column=6,
                 value=f'=BDP({bbg_id},"CUR_MKT_CAP")').font = FORMULA_FONT

        if needs_adv:
            ws.cell(row=row, column=7,
                     value=f'=BDP({bbg_id},"VOLUME_AVG_20D")').font = FORMULA_FONT
            ws.cell(row=row, column=8,
                     value=f'=BDP({bbg_id},"VOLUME_AVG_3M")').font = FORMULA_FONT
        else:
            note = ws.cell(row=row, column=7, value="N/A -- warrant, no ADV model")
            note.font = LABEL_FONT

        # Fallback ticker-resolution source for security_master.py -- pulled
        # for every security (including warrants, unlike ADV) since it's
        # cheap and useful as a human-readable sanity check on its own.
        ws.cell(row=row, column=9,
                 value=f'=BDP({bbg_id},"PARSEKYABLE_DES")').font = FORMULA_FONT

        # Sector classification for analyze.py's compute_sector_concentration.
        # Same reasoning as PARSEKYABLE_DES: pulled for every security
        # (including warrants) since a warrant's underlying company still
        # has a sector, even though the warrant itself has no ADV model.
        ws.cell(row=row, column=10,
                 value=f'=BDP({bbg_id},"GICS_INDUSTRY_NAME")').font = FORMULA_FONT
        ws.cell(row=row, column=11,
                 value=f'=BDP({bbg_id},"GICS_SUB_INDUSTRY_NAME")').font = FORMULA_FONT

    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["C"].width = 24

    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    adv_needed = sum(1 for s in securities if not s["classes"].issubset(NO_ADV_CLASSES))
    print(f"Wrote {output_path}")
    print(f"{len(securities)} distinct CUSIPs ({len(rows)} raw filing rows collapsed to this)")
    print(f"{adv_needed} need ADV fields; {len(securities) - adv_needed} are warrant-only (price only)")
    print("\nDO NOT run recalc.py on this file -- see module docstring.")
    print("PX_LAST, EQY_SH_OUT, CUR_MKT_CAP, VOLUME_AVG_20D, VOLUME_AVG_3M, and the /cusip/ "
          "identifier syntax are user-confirmed against a live Bloomberg session.")
    print("GICS_INDUSTRY_NAME, GICS_SUB_INDUSTRY_NAME: mnemonics live-confirmed 2026-09-08 "
          "(see SKILL.md) -- not yet confirmed pulled through THIS per-CUSIP template at scale; "
          "check these two columns specifically on the first real refresh of this file.")
    return output_path


def main(argv=None):
    argv = argv if argv is not None else sys.argv
    input_file = DEFAULT_INPUT_FILE
    output_file = argv[1] if len(argv) > 1 else DEFAULT_OUTPUT_FILE
    with open(input_file) as f:
        rows = json.load(f)
    export_template(rows, output_file)


if __name__ == "__main__":
    main()
